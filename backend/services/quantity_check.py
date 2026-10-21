"""Compare calculation-sheet totals with the quantities stored on BOQ items."""

import logging
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from utils.calculation_sheet_utils import FIRST_ITEM_COLUMN_INDEX

TOLERANCE = 0.01

# Row 2 of an item column is its invoice number. Column C of that row is the
# sheet invoice, used only when the sheet has a single item.
# Row 5 is the section number. Row 24 is Quantity Submitted.
EXCEL_INVOICE_ROW = 2
SHEET_INVOICE_COLUMN_INDEX = 2
EXCEL_SECTION_ROW = 5
EXCEL_SUBMITTED_ROW = 24

SheetTotals = Dict[str, Tuple[float, float]]
logger = logging.getLogger(__name__)


def _round3(value: float) -> float:
    return round(float(value or 0), 3)


def quantity_mismatches(
    sheet_totals: SheetTotals,
    boq_totals: SheetTotals,
    manual_sections: Optional[Iterable[str]] = None,
) -> List[dict]:
    """Rows where sheet calculated or submitted totals differ from the BOQ.

    Each map is section_number -> (calculated quantity, submitted quantity).
    Difference is sheet minus system, rounded to 3 decimals.
    A row is a mismatch when either absolute difference is at least 0.01.
    """
    manual = {section for section in (manual_sections or []) if section}
    rows = []
    for section in sorted(set(sheet_totals) | set(boq_totals)):
        sheet_calculated, sheet_submitted = sheet_totals.get(section, (0.0, 0.0))
        system_calculated, system_submitted = boq_totals.get(section, (0.0, 0.0))
        calculated_diff = _round3(sheet_calculated - system_calculated)
        submitted_diff = _round3(sheet_submitted - system_submitted)
        if abs(calculated_diff) < TOLERANCE and abs(submitted_diff) < TOLERANCE:
            continue
        rows.append(
            {
                "section_number": section,
                "sheet_calculated": _round3(sheet_calculated),
                "system_calculated": _round3(system_calculated),
                "calculated_diff": calculated_diff,
                "sheet_submitted": _round3(sheet_submitted),
                "system_submitted": _round3(system_submitted),
                "submitted_diff": submitted_diff,
                "status": "שגיאה/חריגה",
                "has_manual_entries": section in manual,
                "is_non_boq": section not in boq_totals,
            }
        )
    return rows


def _section_text(value) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, float) and value != value:
        return None
    text = str(value).strip()
    if not text or text.lower() in {"nan", "none", "<na>", "nat"}:
        return None
    return text


def submitted_totals_from_rows(
    section_row: Optional[Sequence],
    submitted_row: Optional[Sequence],
    invoice_row: Optional[Sequence] = None,
) -> Dict[str, float]:
    """Sum Excel row 24 for item columns that have an invoice number.

    A column with an empty invoice was not submitted, so its row 24 is left
    out. On a sheet with one item, an empty item invoice falls back to the
    sheet invoice in column C. The same section is added together.
    """
    from utils.quantity_utils import round_quantity

    section_row = section_row or ()
    submitted_row = submitted_row or ()
    invoice_row = invoice_row or ()
    width = max(len(section_row), len(submitted_row), len(invoice_row))
    item_columns = []
    for col in range(FIRST_ITEM_COLUMN_INDEX, width):
        section = _section_text(section_row[col] if col < len(section_row) else None)
        if section:
            item_columns.append((col, section))

    sheet_invoice = _section_text(
        invoice_row[SHEET_INVOICE_COLUMN_INDEX]
        if SHEET_INVOICE_COLUMN_INDEX < len(invoice_row)
        else None
    )
    single_item = len(item_columns) == 1
    totals: Dict[str, float] = {}
    for col, section in item_columns:
        invoice = _section_text(invoice_row[col] if col < len(invoice_row) else None)
        if not invoice and single_item:
            invoice = sheet_invoice
        if not invoice:
            continue
        quantity = round_quantity(
            submitted_row[col] if col < len(submitted_row) else 0
        )
        totals[section] = totals.get(section, 0.0) + quantity
    return totals


def submitted_by_section_from_workbook(path: str) -> Dict[str, float]:
    """Read each item's Quantity Submitted cell (Excel row 24) from a calc sheet."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Calculation" not in workbook.sheetnames:
            return {}
        worksheet = workbook["Calculation"]
        max_column = worksheet.max_column or 0
        invoice_row = None
        section_row = None
        submitted_row = None
        for index, row in enumerate(
            worksheet.iter_rows(
                min_row=EXCEL_INVOICE_ROW,
                max_row=EXCEL_SUBMITTED_ROW,
                max_col=max_column or None,
                values_only=True,
            ),
            start=EXCEL_INVOICE_ROW,
        ):
            if index == EXCEL_INVOICE_ROW:
                invoice_row = row
            elif index == EXCEL_SECTION_ROW:
                section_row = row
            elif index == EXCEL_SUBMITTED_ROW:
                submitted_row = row
        return submitted_totals_from_rows(section_row, submitted_row, invoice_row)
    finally:
        workbook.close()


def mismatches_from_db(db) -> List[dict]:
    from pathlib import Path

    from models import models
    from utils.concentration_utils import entry_cumulative_submitted

    calculated_by_section: Dict[str, float] = {}
    for entry in db.query(models.CalculationEntry).all():
        section = entry.section_number
        if not section:
            continue
        calculated_by_section[section] = calculated_by_section.get(
            section, 0.0
        ) + float(entry.estimated_quantity or 0)

    submitted_by_section: Dict[str, float] = {}
    for sheet in db.query(models.CalculationSheet).all():
        path = sheet.source_file_path
        if not path or not Path(path).is_file():
            continue
        try:
            file_totals = submitted_by_section_from_workbook(path)
        except Exception:
            logger.warning(
                "Could not read quantity submitted (row 24) from %s",
                path,
                exc_info=True,
            )
            continue
        for section, quantity in file_totals.items():
            submitted_by_section[section] = (
                submitted_by_section.get(section, 0.0) + quantity
            )

    sheet_totals: SheetTotals = {}
    for section in set(calculated_by_section) | set(submitted_by_section):
        sheet_totals[section] = (
            calculated_by_section.get(section, 0.0),
            submitted_by_section.get(section, 0.0),
        )

    concentration_submitted: Dict[str, float] = {}
    sections_with_entries = set()
    concentration_rows = (
        db.query(
            models.BOQItem.section_number,
            models.ConcentrationEntry,
        )
        .join(
            models.ConcentrationSheet,
            models.ConcentrationSheet.boq_item_id == models.BOQItem.id,
        )
        .join(
            models.ConcentrationEntry,
            models.ConcentrationEntry.concentration_sheet_id
            == models.ConcentrationSheet.id,
        )
        .all()
    )
    for section, entry in concentration_rows:
        if not section:
            continue
        sections_with_entries.add(section)
        concentration_submitted[section] = concentration_submitted.get(
            section, 0.0
        ) + entry_cumulative_submitted(entry)

    boq_totals: SheetTotals = {}
    for item in db.query(models.BOQItem).all():
        if item.section_number in sections_with_entries:
            submitted = concentration_submitted.get(item.section_number, 0.0)
        else:
            submitted = float(item.quantity_submitted or 0)
        boq_totals[item.section_number] = (
            float(item.estimated_quantity or 0),
            submitted,
        )

    manual_sections = {
        section
        for (section,) in db.query(models.BOQItem.section_number)
        .join(
            models.ConcentrationSheet,
            models.ConcentrationSheet.boq_item_id == models.BOQItem.id,
        )
        .join(
            models.ConcentrationEntry,
            models.ConcentrationEntry.concentration_sheet_id == models.ConcentrationSheet.id,
        )
        .filter(models.ConcentrationEntry.is_manual == True)
        .distinct()
        .all()
        if section
    }
    return quantity_mismatches(sheet_totals, boq_totals, manual_sections)
