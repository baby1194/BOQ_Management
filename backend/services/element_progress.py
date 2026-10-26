"""Which contract items for one concrete element have been submitted."""

import re
from typing import Iterable, List

_ELEMENT_CODE = re.compile(r"(?<![\w.])([A-Za-z]{1,5})\s?-\s?(\d{1,4}[A-Za-z]?)(?!\w)")
_SHORT_NAME = re.compile(r"^(?=.*[^\W\d_])(?=.*\d)[\w\- /]{2,20}$")
_DASH = re.compile(r"\s*-\s*")
_WORD_SPLIT = re.compile(r"[\n,;/|]+")


def _normalize(text) -> str:
    return _DASH.sub("-", str(text or "")).strip().casefold()


def element_names_in(text, whole_value: bool = False) -> List[str]:
    """Element codes such as DC-36 inside a description. whole_value also accepts a short name like 'Wall 4'."""
    value = str(text or "").strip()
    names = [f"{letters.upper()}-{number.upper()}" for letters, number in _ELEMENT_CODE.findall(value)]
    if not names and whole_value and _SHORT_NAME.match(value) and not re.fullmatch(r"[\d.\-/ ]+", value):
        names.append(_DASH.sub("-", value))
    return names


def element_names_from_entries(entries: Iterable[dict]) -> List[str]:
    """Names found in concentration Description and Work Description."""
    found = {}
    for entry in entries:
        names = element_names_in(entry.get("description"), whole_value=True)
        names += element_names_in(entry.get("invoice_description"))
        for name in names:
            found.setdefault(name.casefold(), name)
    return sorted(found.values(), key=str.casefold)


def split_search_words(text: str) -> List[str]:
    """Words typed as RW-18 / DC-50 / rails, in the order they were written."""
    words = []
    seen = set()
    for part in _WORD_SPLIT.split(str(text or "")):
        word = part.strip()
        key = _normalize(word)
        if not key or key in seen:
            continue
        seen.add(key)
        words.append(word)
    return words


def mentions(text, element: str) -> bool:
    needle = _normalize(element)
    if not needle:
        return False
    return re.search(rf"(?<!\w){re.escape(needle)}(?!\w)", _normalize(text)) is not None


def items_for_element(
    element: str,
    boq_items: Iterable[dict],
    entries_by_section: dict,
    drawing_names: Iterable[str] = (),
) -> List[dict]:
    """Contract items tied to the element, with the quantities of the rows that name it."""
    if not (element or "").strip():
        return []
    names = [element] + [name for name in drawing_names if name and str(name).strip()]

    def linked(text) -> bool:
        return any(mentions(text, name) for name in names)

    rows = []
    for item in boq_items:
        section = item.get("section_number") or ""
        entries = [
            entry for entry in entries_by_section.get(section, [])
            if linked(entry.get("description"))
            or linked(entry.get("invoice_description"))
            or linked(entry.get("drawing_no"))
            or linked(entry.get("calculation_sheet_no"))
        ]
        if not entries and not linked(item.get("description")):
            continue
        if entries:
            estimated = sum(float(e.get("estimated_quantity") or 0) for e in entries)
            submitted = sum(float(e.get("quantity_submitted") or 0) for e in entries)
        else:
            estimated = float(item.get("estimated_quantity") or 0)
            submitted = float(item.get("quantity_submitted") or 0)
        rows.append(
            {
                "section_number": section,
                "description": item.get("description") or "",
                "contract_quantity": float(item.get("contract_quantity") or 0),
                "estimated_quantity": round(estimated, 3),
                "quantity_submitted": round(submitted, 3),
                "submitted": submitted > 0,
                "calculation_sheet_nos": sorted(
                    {str(e.get("calculation_sheet_no")) for e in entries if e.get("calculation_sheet_no")}
                ),
            }
        )
    rows.sort(key=lambda row: (not row["submitted"], row["section_number"]))
    return rows


# Same figures as one BOQ table row, with the searched word in front.
BOQ_TABLE_KEYS = (
    "serial_number",
    "structure",
    "system",
    "section_number",
    "description",
    "unit",
    "original_contract_quantity",
    "price",
    "submitted_price",
    "submitted_to_supervisor",
    "submitted_to_supervisor_date",
    "total_contract_sum",
    "estimated_quantity",
    "quantity_submitted",
    "internal_quantity",
    "approved_by_project_manager",
    "approved_signed_quantity",
    "partially_submitted_quantity",
    "quantity_decrease",
    "quantity_increase",
    "submitted_increase_quantity",
    "total_estimate",
    "total_submitted",
    "internal_total",
    "total_approved_by_project_manager",
    "approved_signed_total",
    "partial_submitted_total",
    "total_decrease",
    "total_increase",
    "subsection",
    "notes",
)

_WORD_HEADER = {
    "en": {
        "word": "Word",
        "submitted_price": "Submitted price",
        "submitted_to_supervisor": "Submitted Increase/Decrease",
        "submitted_to_supervisor_date": "Submitted date Increase/Decrease",
        "quantity_decrease": "Quantity Decrease",
        "quantity_increase": "Quantity Increase",
        "submitted_increase_quantity": "Submitted increasing quantity",
    },
    "he": {
        "word": "מילה",
        "submitted_price": "מחיר שהוגש",
        "submitted_to_supervisor": "הוגש הגדלה/הקטנה",
        "submitted_to_supervisor_date": "תאריך הגשה הגדלה/הקטנה",
        "quantity_decrease": "הקטנות",
        "quantity_increase": "הגדלות",
        "submitted_increase_quantity": "הוגשה הגדלה",
    },
}

_TEXT_KEYS = {
    "word",
    "section_number",
    "description",
    "unit",
    "system",
    "subsection",
    "notes",
    "submitted_to_supervisor",
    "submitted_to_supervisor_date",
}


def _number(value, default=0.0) -> float:
    if value is None or value == "":
        return default
    return float(value)


def boq_table_fields(item: dict, contract_updates: Iterable[dict] = ()) -> dict:
    """BOQ table columns for one item. contract_updates are {id, quantity, sum} in order."""
    price = _number(item.get("price"))
    estimated = _number(item.get("estimated_quantity"))
    submitted = _number(item.get("quantity_submitted"))
    approved_signed = _number(item.get("approved_signed_quantity"))
    original = item.get("original_contract_quantity")
    if original is None:
        original = item.get("contract_quantity")
    original = _number(original)
    current = original
    fields = {
        "serial_number": item.get("serial_number") if item.get("serial_number") is not None else "",
        "structure": item.get("structure") if item.get("structure") is not None else "",
        "system": item.get("system") or "",
        "section_number": item.get("section_number") or "",
        "description": item.get("description") or "",
        "unit": item.get("unit") or "",
        "original_contract_quantity": original,
        "price": price,
        "submitted_price": _number(item.get("submitted_price")) if item.get("submitted_price") is not None else "",
        "submitted_to_supervisor": bool(item.get("submitted_to_supervisor")),
        "submitted_to_supervisor_date": item.get("submitted_to_supervisor_date") or "",
        "total_contract_sum": _number(item.get("total_contract_sum"), original * price),
        "estimated_quantity": estimated,
        "quantity_submitted": submitted,
        "internal_quantity": _number(item.get("internal_quantity")),
        "approved_by_project_manager": _number(item.get("approved_by_project_manager")),
        "approved_signed_quantity": approved_signed,
        "submitted_increase_quantity": _number(item.get("submitted_increase_quantity")),
        "total_estimate": _number(item.get("total_estimate"), estimated * price),
        "total_submitted": _number(item.get("total_submitted"), submitted * price),
        "internal_total": _number(item.get("internal_total")),
        "total_approved_by_project_manager": _number(item.get("total_approved_by_project_manager")),
        "approved_signed_total": _number(item.get("approved_signed_total"), approved_signed * price),
        "subsection": item.get("subsection") or "",
        "notes": item.get("notes") or "",
        "contract_quantity": original,
    }
    for update in contract_updates or ():
        quantity = update.get("quantity")
        total = update.get("sum")
        fields[f"updated_contract_quantity_{update['id']}"] = "" if quantity is None else _number(quantity)
        fields[f"updated_contract_sum_{update['id']}"] = "" if total is None else _number(total)
        if quantity is not None:
            current = _number(quantity)
            fields["contract_quantity"] = current
    decrease = max(0.0, current - estimated)
    increase = max(0.0, estimated - current)
    partial_qty = submitted - approved_signed
    fields["quantity_decrease"] = decrease
    fields["quantity_increase"] = increase
    fields["partially_submitted_quantity"] = partial_qty
    fields["partial_submitted_total"] = partial_qty * price
    fields["total_decrease"] = decrease * price
    fields["total_increase"] = increase * price
    return fields


def _blank_boq_row() -> dict:
    blank = {key: "" for key in BOQ_TABLE_KEYS}
    blank["submitted_to_supervisor"] = ""
    return blank


def rows_for_words(
    text: str,
    boq_items: Iterable[dict],
    entries_by_section: dict,
    drawing_names_by_element: dict = None,
) -> List[dict]:
    """One BOQ row per word and connected section. A word with no items stays visible."""
    drawing_names_by_element = drawing_names_by_element or {}
    boq_items = [
        item if "partially_submitted_quantity" in item else boq_table_fields(item)
        for item in boq_items
    ]
    by_section = {item.get("section_number") or "": item for item in boq_items}
    table = []
    for word in split_search_words(text):
        names = drawing_names_by_element.get(_normalize(word), [])
        found = items_for_element(word, boq_items, entries_by_section, names)
        known = {row["section_number"] for row in found}
        needle = _normalize(word)
        for item in boq_items:
            section = item.get("section_number") or ""
            if section in known or needle not in _normalize(section):
                continue
            found.append({"section_number": section})
            known.add(section)
        if not found:
            table.append({"word": word, **_blank_boq_row()})
            continue
        for row in found:
            item = by_section.get(row["section_number"], {})
            table.append({"word": word, **_blank_boq_row(), **{
                key: item.get(key, "")
                for key in list(BOQ_TABLE_KEYS) + [
                    key for key in item if str(key).startswith("updated_contract_")
                ]
            }})
    return table


def word_items_workbook(rows: List[dict], language: str = "en", contract_updates: Iterable = ()) -> bytes:
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from services.excel_service import _boq_excel_number_format, _boq_header_label

    hebrew = language == "he"
    lang = "he" if hebrew else "en"
    labels = _WORD_HEADER[lang]
    yes = "כן" if hebrew else "Yes"
    no = "לא" if hebrew else "No"
    columns = ["word"]
    columns.extend(BOQ_TABLE_KEYS[: BOQ_TABLE_KEYS.index("total_contract_sum") + 1])
    updates = list(contract_updates or ())
    for update in updates:
        columns.append(f"updated_contract_quantity_{update.id}")
    for update in updates:
        columns.append(f"updated_contract_sum_{update.id}")
    columns.extend(BOQ_TABLE_KEYS[BOQ_TABLE_KEYS.index("estimated_quantity"):])

    def header(key: str) -> str:
        if key in labels:
            return labels[key]
        return _boq_header_label(key, lang, {update.id: update for update in updates})

    book = Workbook()
    sheet = book.active
    sheet.title = "Items"
    sheet.sheet_view.rightToLeft = hebrew
    sheet.append([header(key) for key in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        values = []
        for key in columns:
            value = row.get(key, "")
            if key == "submitted_to_supervisor" and value != "":
                value = yes if value else no
            values.append(value)
        sheet.append(values)
        excel_row = sheet.max_row
        for index, key in enumerate(columns, start=1):
            if key in _TEXT_KEYS:
                continue
            cell = sheet.cell(excel_row, index)
            if cell.value == "" or cell.value is None:
                continue
            cell.number_format = _boq_excel_number_format(key)
    for index, key in enumerate(columns, start=1):
        width = 42 if key == "description" else 18 if key == "word" else 16
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "B2"
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
