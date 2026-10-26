"""Write a PDF copy of each calculation-sheet workbook into its own folder."""

import os
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Iterable, List, Optional, Tuple

from bidi.algorithm import get_display
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.lib.utils import simpleSplit
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Table, TableStyle

CALCULATION_SHEET_NAME = "Calculation"
_FONT_NAME = "SheetPrintHebrew"
_DEFAULT_COLUMN_WIDTH = 8.43
_MIN_FONT_SIZE = 6.5
# The title block of a calculation sheet. Repeated on every page, as in Excel.
HEADER_ROW_COUNT = 27


@dataclass
class PrintGrid:
    rows: List[List[str]]
    widths: List[float]
    right_to_left: bool
    spans: List[Tuple[Tuple[int, int], Tuple[int, int]]] = field(default_factory=list)
    numeric: List[Tuple[int, int]] = field(default_factory=list)
    repeat_rows: int = 0


def pdf_beside_workbook(workbook: Path) -> Path:
    return workbook.with_suffix(".pdf")


def _has_hebrew(text: str) -> bool:
    return any("\u0590" <= char <= "\u05ff" for char in text)


def format_cell(value) -> str:
    """Show values the way a person reads them: no float noise, dates as day/month/year."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float):
        rounded = round(value, 3)
        if rounded == int(rounded):
            return str(int(rounded))
        return f"{rounded:.3f}".rstrip("0").rstrip(".")
    return str(value).strip()


def _meaningful(value) -> bool:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return round(value, 6) != 0
    return format_cell(value) != ""


def calculation_worksheet(workbook) -> Worksheet:
    if CALCULATION_SHEET_NAME in workbook.sheetnames:
        return workbook[CALCULATION_SHEET_NAME]
    return workbook.worksheets[0]


def _filled_bounds(ws: Worksheet) -> Optional[Tuple[int, int, int, int]]:
    """Print area when one is set, otherwise the smallest block holding every value."""
    area = ws.print_area
    if area:
        first = str(area).split(",")[0]
        ref = first.split("!")[-1].replace("$", "").replace("'", "")
        min_col, min_row, max_col, max_row = range_boundaries(ref)
        return min_row, min_col, max_row, max_col
    bounds = None
    for row in ws.iter_rows():
        for cell in row:
            if format_cell(cell.value) == "":
                continue
            if bounds is None:
                bounds = [cell.row, cell.column, cell.row, cell.column]
            else:
                bounds[0] = min(bounds[0], cell.row)
                bounds[1] = min(bounds[1], cell.column)
                bounds[2] = max(bounds[2], cell.row)
                bounds[3] = max(bounds[3], cell.column)
    return tuple(bounds) if bounds else None


def _row_hidden(ws: Worksheet, row: int) -> bool:
    dimension = ws.row_dimensions.get(row)
    return bool(dimension is not None and dimension.hidden)


def _hidden_columns(ws: Worksheet) -> set:
    hidden = set()
    for dimension in ws.column_dimensions.values():
        if dimension.hidden and dimension.min and dimension.max:
            hidden.update(range(dimension.min, dimension.max + 1))
    return hidden


def _column_width(ws: Worksheet, column: int) -> float:
    from openpyxl.utils import get_column_letter

    dimension = ws.column_dimensions.get(get_column_letter(column))
    if dimension is not None and dimension.width:
        return float(dimension.width)
    return _DEFAULT_COLUMN_WIDTH


def filled_cell_boxes(
    rows: List[List[str]],
    spans: List[Tuple[Tuple[int, int], Tuple[int, int]]],
) -> List[Tuple[Tuple[int, int], Tuple[int, int]]]:
    """Border each filled cell, and each merged block that holds a value."""
    covered = set()
    boxes = []
    for (c0, r0), (c1, r1) in spans:
        cells = [(c, r) for r in range(r0, r1 + 1) for c in range(c0, c1 + 1)]
        covered.update(cells)
        if any(rows[r][c] for c, r in cells):
            boxes.append(((c0, r0), (c1, r1)))
    for ri, row in enumerate(rows):
        for ci, text in enumerate(row):
            if text and (ci, ri) not in covered:
                boxes.append(((ci, ri), (ci, ri)))
    return boxes


def build_print_grid(ws: Worksheet, columns: Optional[Iterable[int]] = None) -> PrintGrid:
    """Filled columns of the sheet, laid out like Excel.

    Columns come from the sheet itself. When the sheet continues past the title
    block, every one of those title rows is kept so it can repeat on each page.
    """
    bounds = _filled_bounds(ws)
    if bounds is None:
        return PrintGrid(rows=[[""]], widths=[1.0], right_to_left=False)
    min_row, min_col, max_row, max_col = bounds

    chosen = {int(col) for col in (columns or []) if int(col) >= 1}
    hidden = _hidden_columns(ws)
    column_range = [
        c for c in range(min_col, max_col + 1)
        if (not chosen or c in chosen) and c not in hidden
    ]

    header_rows = (
        [r for r in range(1, HEADER_ROW_COUNT + 1) if not _row_hidden(ws, r)]
        if max_row > HEADER_ROW_COUNT
        else []
    )
    body_rows = [
        r for r in range(min_row, max_row + 1)
        if r > HEADER_ROW_COUNT and not _row_hidden(ws, r)
    ]
    if not header_rows:
        body_rows = [r for r in range(min_row, max_row + 1) if not _row_hidden(ws, r)]

    rows_to_read = header_rows + [r for r in body_rows if r not in header_rows]
    values = {
        (r, c): ws.cell(row=r, column=c).value
        for r in rows_to_read
        for c in column_range
    }
    # Formula columns that only hold 0 are part of Excel's used range but not of the sheet.
    kept_cols = [
        c for c in column_range
        if any(_meaningful(values[(r, c)]) for r in rows_to_read)
    ]
    kept_body = [
        r for r in body_rows
        if any(format_cell(values[(r, c)]) for c in kept_cols)
    ]
    kept_rows = header_rows + [r for r in kept_body if r not in header_rows] if header_rows else kept_body
    if not kept_rows or not kept_cols:
        return PrintGrid(rows=[[""]], widths=[1.0], right_to_left=False)
    repeat_rows = len(header_rows) if kept_body else 0

    texts = [[format_cell(values[(r, c)]) for c in kept_cols] for r in kept_rows]
    right_to_left = bool(ws.sheet_view.rightToLeft) or any(
        _has_hebrew(text) for row in texts for text in row
    )

    row_pos = {r: i for i, r in enumerate(kept_rows)}
    col_pos = {c: i for i, c in enumerate(kept_cols)}
    count = len(kept_cols)

    def place(index: int) -> int:
        return count - 1 - index if right_to_left else index

    spans = []
    for merged in ws.merged_cells.ranges:
        rows_in = [row_pos[r] for r in range(merged.min_row, merged.max_row + 1) if r in row_pos]
        cols_in = [col_pos[c] for c in range(merged.min_col, merged.max_col + 1) if c in col_pos]
        if (merged.min_row not in row_pos or merged.min_col not in col_pos
                or (len(rows_in) < 2 and len(cols_in) < 2)):
            continue
        a, b = place(min(cols_in)), place(max(cols_in))
        spans.append(((min(a, b), min(rows_in)), (max(a, b), max(rows_in))))

    numeric = [
        (place(ci), ri)
        for ri, r in enumerate(kept_rows)
        for ci, c in enumerate(kept_cols)
        if isinstance(values[(r, c)], (int, float)) and not isinstance(values[(r, c)], bool)
    ]

    widths = [_column_width(ws, c) for c in kept_cols]
    shown = [[_display(text) for text in row] for row in texts]
    if right_to_left:
        shown = [list(reversed(row)) for row in shown]
        widths = list(reversed(widths))
    for start, end in spans:
        _move_text_to_span_start(shown, start, end)
    return PrintGrid(
        rows=shown,
        widths=widths,
        right_to_left=right_to_left,
        spans=spans,
        numeric=numeric,
        repeat_rows=repeat_rows,
    )


def _move_text_to_span_start(rows: List[List[str]], start, end) -> None:
    """ReportLab draws a span from its first cell, so the text has to live there."""
    (c0, r0), (c1, r1) = start, end
    if rows[r0][c0]:
        return
    for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            if rows[r][c]:
                rows[r0][c0], rows[r][c] = rows[r][c], ""
                return


def overflow_spans(
    grid: PrintGrid,
    col_widths: List[float],
    font_name: str,
    font_size: float,
    padding: float,
) -> Tuple[list, list]:
    """Long text runs into empty neighbour cells, as in Excel. Returns spans and per-cell font sizes."""
    covered = set()
    merged_width = {}
    for (c0, r0), (c1, r1) in grid.spans:
        covered.update((c, r) for c in range(c0, c1 + 1) for r in range(r0, r1 + 1))
        merged_width[(c0, r0)] = sum(col_widths[c0:c1 + 1])
    numeric = set(grid.numeric)
    primary = -1 if grid.right_to_left else 1
    spans, sizes = [], []
    count = len(col_widths)

    def fit(cell, text, have):
        need = pdfmetrics.stringWidth(text, font_name, font_size) + 2 * padding
        if have >= need:
            return
        room = max(have - 2 * padding, 1.0)
        size = font_size * room / (need - 2 * padding)
        if size >= _MIN_FONT_SIZE or cell in numeric:
            sizes.append((cell, max(4.0, size)))
            return
        lines = simpleSplit(text, font_name, _MIN_FONT_SIZE, room)
        if _has_hebrew(text):
            # The text is already in visual order, so its last words come first.
            lines.reverse()
        col, row_index = cell
        grid.rows[row_index][col] = "\n".join(lines)
        sizes.append((cell, _MIN_FONT_SIZE))

    for ri, row in enumerate(grid.rows):
        for ci, text in enumerate(row):
            if not text:
                continue
            if (ci, ri) in merged_width:
                fit((ci, ri), text, merged_width[(ci, ri)])
                continue
            if (ci, ri) in covered:
                continue
            need = pdfmetrics.stringWidth(text, font_name, font_size) + 2 * padding
            first = last = ci
            have = col_widths[ci]
            if (ci, ri) not in numeric:
                for step in (primary, -primary):
                    edge = ci
                    while have < need:
                        nxt = edge + step
                        if not 0 <= nxt < count or row[nxt] or (nxt, ri) in covered:
                            break
                        edge = nxt
                        have += col_widths[nxt]
                    first, last = min(first, edge), max(last, edge)
            if first != last:
                spans.append(((first, ri), (last, ri)))
                covered.update((c, ri) for c in range(first, last + 1))
                row[ci] = ""
                row[first] = text
            fit((first, ri), text, have)
    return spans, sizes


def _display(text: str) -> str:
    return get_display(text) if _has_hebrew(text) else text


def _ensure_font() -> str:
    if _FONT_NAME in pdfmetrics.getRegisteredFontNames():
        return _FONT_NAME
    for path in (
        r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\ARIAL.TTF",
        r"C:\Windows\Fonts\tahoma.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ):
        if os.path.exists(path):
            pdfmetrics.registerFont(TTFont(_FONT_NAME, path))
            return _FONT_NAME
    return "Helvetica"


def write_sheet_pdf(
    workbook: Path,
    destination: Path,
    orientation: str = "landscape",
    margin_mm: float = 10,
    columns: Optional[Iterable[int]] = None,
) -> None:
    book = load_workbook(workbook, data_only=True)
    try:
        grid = build_print_grid(calculation_worksheet(book), columns)
    finally:
        book.close()

    page = landscape(A4) if orientation == "landscape" else A4
    margin = max(0.0, float(margin_mm)) * mm
    document = BaseDocTemplate(str(destination), pagesize=page)
    frame = Frame(
        margin,
        margin,
        page[0] - 2 * margin,
        page[1] - 2 * margin,
        leftPadding=0,
        rightPadding=0,
        topPadding=0,
        bottomPadding=0,
    )
    document.addPageTemplates([PageTemplate(frames=[frame])])
    usable = page[0] - (2 * margin)
    total = sum(grid.widths) or 1.0
    col_widths = [usable * width / total for width in grid.widths]
    font_name = _ensure_font()
    font_size = 8
    padding = 2
    extra_spans, sizes = overflow_spans(grid, col_widths, font_name, font_size, padding)
    align = "RIGHT" if grid.right_to_left else "LEFT"
    style = [
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("FONTSIZE", (0, 0), (-1, -1), font_size),
        ("LEADING", (0, 0), (-1, -1), font_size * 1.2),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, -1), align),
        ("LEFTPADDING", (0, 0), (-1, -1), padding),
        ("RIGHTPADDING", (0, 0), (-1, -1), padding),
    ]
    style += [("SPAN", start, end) for start, end in grid.spans + extra_spans]
    style += [
        ("BOX", start, end, 0.25, colors.grey)
        for start, end in filled_cell_boxes(grid.rows, grid.spans)
    ]
    style += [("ALIGN", cell, cell, "CENTER") for cell in grid.numeric]
    style += [("FONTSIZE", cell, cell, size) for cell, size in sizes]
    style += [("LEADING", cell, cell, size * 1.2) for cell, size in sizes]
    table = Table(
        grid.rows,
        colWidths=col_widths,
        repeatRows=grid.repeat_rows,
        hAlign="RIGHT" if grid.right_to_left else "LEFT",
    )
    table.setStyle(TableStyle(style))
    document.build([table])


def print_workbooks(
    roots: Iterable[Path],
    orientation: str = "landscape",
    margin_mm: float = 10,
    columns: Optional[List[int]] = None,
) -> List[str]:
    written: List[str] = []
    for root in roots:
        if not root.exists():
            continue
        for workbook in root.rglob("*.xlsx"):
            if workbook.name.startswith("~$"):
                continue
            destination = pdf_beside_workbook(workbook)
            write_sheet_pdf(
                workbook,
                destination,
                orientation=orientation,
                margin_mm=margin_mm,
                columns=columns,
            )
            written.append(str(destination))
    return written
