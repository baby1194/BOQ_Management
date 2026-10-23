import pandas as pd
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from models import models
from services.non_boq_service import (
    build_non_boq_import_rows,
    mark_non_boq_items_moved,
    non_boq_import_rows,
    select_import_rows,
)
from utils.calculation_sheet_utils import read_entry_system, read_non_boq_price


def _session():
    engine = create_engine("sqlite://")
    models.Base.metadata.create_all(bind=engine)
    return sessionmaker(bind=engine)()


def test_new_project_remembers_the_main_project(tmp_path, monkeypatch):
    from database import project_registry

    monkeypatch.setattr(project_registry, "DATABASE_DIR", tmp_path)
    monkeypatch.setattr(project_registry, "REGISTRY_FILE", tmp_path / "projects.json")
    linked = project_registry.create_project_entry("Extras", source_project_id="main-1")
    plain = project_registry.create_project_entry("Plain")
    saved = {p["id"]: p for p in project_registry.list_projects()}
    assert saved[linked["id"]]["source_project_id"] == "main-1"
    assert "source_project_id" not in saved[plain["id"]]


def test_select_import_rows_keeps_all_when_nothing_chosen():
    rows = [{"section_number": "A"}, {"section_number": "B"}]
    assert select_import_rows(rows, None) == rows
    assert select_import_rows(rows, [" B "]) == [{"section_number": "B"}]
    assert select_import_rows(rows, []) == []


def test_moved_items_record_the_target_project():
    db = _session()
    db.add_all([models.NonBoqItem(section_number="A"), models.NonBoqItem(section_number="B")])
    db.commit()
    assert mark_non_boq_items_moved(db, ["A"], "project-2") == 1
    db.commit()
    moved = {item.section_number: item.moved_to_project_id for item in db.query(models.NonBoqItem)}
    assert moved == {"A": "project-2", "B": None}


def test_price_is_read_from_excel_row_7():
    df = pd.DataFrame([[None] * 6 for _ in range(8)])
    df.iloc[6, 4] = 1213
    assert read_non_boq_price(df, 4) == 1213


def test_system_is_read_from_excel_row_4():
    df = pd.DataFrame([[None] * 6 for _ in range(8)])
    df.iloc[3, 4] = 919.0
    assert read_entry_system(df, 4) == "919"


def test_import_rows_sum_quantity_and_keep_the_first_price():
    rows = build_non_boq_import_rows(
        [
            {
                "section_number": "99.02.01.0640",
                "description": "Extra wall",
                "estimated_quantity": 10,
                "quantity_submitted": 4,
                "unit_price": 50,
                "system": "919",
            },
            {
                "section_number": "99.02.01.0640",
                "description": "",
                "estimated_quantity": 2.5,
                "unit_price": 80,
            },
            {
                "section_number": "99.05.031.0012",
                "description": "",
                "estimated_quantity": 4,
                "unit_price": None,
            },
        ]
    )
    by_section = {row["section_number"]: row for row in rows}
    assert by_section["99.02.01.0640"]["quantity"] == 12.5
    assert by_section["99.02.01.0640"]["price"] == 50
    assert by_section["99.02.01.0640"]["quantity_submitted"] == 4
    assert by_section["99.02.01.0640"]["system"] == "919"
    assert by_section["99.02.01.0640"]["total"] == 625
    assert by_section["99.05.031.0012"]["description"] == "99.05.031.0012"
    assert by_section["99.05.031.0012"]["missing_price"] is True


def test_import_description_uses_row_22_not_the_invoice_title():
    db = _session()
    db.add(models.NonBoqItem(section_number="99.01.050.0210"))
    sheet = models.CalculationSheet(
        file_name="41-8-1.xlsx",
        calculation_sheet_no="41/8/1",
        drawing_no="07",
        description="sheet header from row 3",
    )
    db.add(sheet)
    db.flush()
    db.add(
        models.CalculationEntry(
            calculation_sheet_id=sheet.id,
            section_number="99.01.050.0210",
            invoice_description="short title from row 3",
            item_description="specification from row 22",
            estimated_quantity=10,
            unit_price=7.92,
        )
    )
    db.commit()

    rows = non_boq_import_rows(db)
    assert rows[0]["description"] == "specification from row 22"


def test_blank_item_description_does_not_use_the_row_3_title():
    db = _session()
    db.add(models.NonBoqItem(section_number="99.01.050.0210"))
    sheet = models.CalculationSheet(
        file_name="41-8-1.xlsx",
        calculation_sheet_no="41/8/1",
        drawing_no="07",
        description="sheet header from row 3",
    )
    db.add(sheet)
    db.flush()
    db.add(
        models.CalculationEntry(
            calculation_sheet_id=sheet.id,
            section_number="99.01.050.0210",
            invoice_description="short title from row 3",
            estimated_quantity=10,
            unit_price=7.92,
        )
    )
    db.commit()

    rows = non_boq_import_rows(db)
    assert rows[0]["description"] == "99.01.050.0210"


def test_saved_sheet_picks_up_row_22_from_the_source_file(tmp_path):
    from openpyxl import Workbook

    book = Workbook()
    book.active.title = "Calculation"
    ws = book.active
    ws["C1"] = "41/8/1"
    ws["C2"] = "07"
    ws["C3"] = "sheet header from row 3"
    ws.cell(row=3, column=5, value="short title from row 3")
    ws.cell(row=5, column=5, value="99.01.050.0210")
    ws.cell(row=6, column=5, value=10)
    ws.cell(row=7, column=5, value=7.92)
    ws.cell(row=22, column=5, value="specification from row 22")
    ws.cell(row=40, column=1, value=" ")
    path = tmp_path / "41-8-1.xlsx"
    book.save(path)

    db = _session()
    db.add(models.NonBoqItem(section_number="99.01.050.0210"))
    sheet = models.CalculationSheet(
        file_name="41-8-1.xlsx",
        calculation_sheet_no="41/8/1",
        drawing_no="07",
        description="sheet header from row 3",
        source_file_path=str(path),
    )
    db.add(sheet)
    db.flush()
    db.add(
        models.CalculationEntry(
            calculation_sheet_id=sheet.id,
            section_number="99.01.050.0210",
            invoice_description="short title from row 3",
            estimated_quantity=10,
            unit_price=7.92,
        )
    )
    db.commit()

    rows = non_boq_import_rows(db)
    assert rows[0]["description"] == "specification from row 22"
    stored = db.query(models.CalculationEntry).one()
    assert stored.item_description == "specification from row 22"
