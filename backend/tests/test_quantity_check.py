from openpyxl import Workbook

from services.quantity_check import (
    submitted_by_section_from_workbook,
    submitted_totals_from_rows,
    quantity_mismatches,
)


def test_matching_totals_are_omitted():
    sheets = {"1.02.01.0736": (590.224, 342.224)}
    boq = {"1.02.01.0736": (590.224, 342.224)}
    assert quantity_mismatches(sheets, boq) == []


def test_reports_calculated_and_submitted_gaps():
    sheets = {"1.1": (10.0, 4.0), "1.2": (3.0, 3.0)}
    boq = {"1.1": (8.0, 4.0), "1.2": (3.0, 1.0)}
    rows = quantity_mismatches(sheets, boq)
    by_section = {row["section_number"]: row for row in rows}
    assert by_section["1.1"]["calculated_diff"] == 2.0
    assert by_section["1.1"]["submitted_diff"] == 0.0
    assert by_section["1.2"]["submitted_diff"] == 2.0
    assert by_section["1.1"]["status"] == "שגיאה/חריגה"


def test_difference_under_one_hundredth_is_ok():
    sheets = {"1.1": (10.004, 5.0)}
    boq = {"1.1": (10.0, 5.0)}
    assert quantity_mismatches(sheets, boq) == []


def test_totals_that_match_on_two_decimals_are_not_listed():
    sheets = {"1.1": (411.944, 184.444)}
    boq = {"1.1": (411.935, 184.435)}
    assert quantity_mismatches(sheets, boq) == []


def test_a_real_cent_of_submitted_difference_is_listed():
    rows = quantity_mismatches({"1.1": (10.0, 184.44)}, {"1.1": (10.0, 184.33)})
    assert rows[0]["submitted_diff"] == 0.11


def test_manual_badge_follows_the_boq_item():
    rows = quantity_mismatches(
        {"1.08.10.3296": (10.0, 1.0)},
        {"1.08.10.3296": (8.0, 1.0)},
        manual_sections={"1.08.10.3296"},
    )
    assert rows[0]["has_manual_entries"] is True


def _grid(height=40, width=14):
    return [[None] * width for _ in range(height)]


def test_submitted_total_sums_detail_lines_that_have_a_prior_invoice():
    rows = _grid()
    rows[1][9] = "07"
    rows[4][9] = "1.02.01.0693"
    rows[23][9] = 63.58
    rows[27][8] = "07"
    rows[27][9] = 36.30
    rows[28][9] = 27.28
    rows[1][11] = "08"
    rows[4][11] = "1.02.01.0010"
    rows[27][10] = "08"
    rows[27][11] = 6.60
    rows[28][11] = 4.96
    assert submitted_totals_from_rows(rows) == {
        "1.02.01.0693": 36.30,
        "1.02.01.0010": 6.60,
    }


def test_item_without_row_2_invoice_is_left_out():
    rows = _grid()
    rows[1][9] = "05"
    rows[4][9] = "1.02.01.0610"
    rows[27][8] = "05"
    rows[27][9] = 77.23
    rows[4][11] = "1.02.01.0641"
    rows[27][10] = "06"
    rows[27][11] = 10
    assert submitted_totals_from_rows(rows) == {"1.02.01.0610": 77.23}


def test_sheet_invoice_in_c2_does_not_count():
    rows = _grid()
    rows[1][2] = "07"
    rows[4][9] = "1.02.01.0610"
    rows[23][9] = 12.5
    rows[27][1] = "07"
    rows[27][9] = 12.5
    assert submitted_totals_from_rows(rows) == {}


def test_workbook_sums_detail_lines_and_ignores_row_24(tmp_path):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Calculation"
    worksheet.cell(row=2, column=10, value="07")
    worksheet.cell(row=5, column=10, value="1.02.01.0693")
    worksheet.cell(row=6, column=10, value=63.58)
    worksheet.cell(row=24, column=10, value=63.58)
    worksheet.cell(row=28, column=9, value="07")
    worksheet.cell(row=28, column=10, value=36.30)
    worksheet.cell(row=29, column=10, value=27.28)
    worksheet.cell(row=2, column=12, value="06")
    worksheet.cell(row=5, column=12, value="1.02.01.0671")
    worksheet.cell(row=24, column=12, value=8.5)
    worksheet.cell(row=28, column=11, value="06")
    worksheet.cell(row=28, column=12, value=4.25)
    worksheet.cell(row=29, column=12, value=9)
    worksheet.cell(row=5, column=14, value="1.02.01.0999")
    worksheet.cell(row=24, column=14, value=50)
    worksheet.cell(row=28, column=13, value="01")
    worksheet.cell(row=28, column=14, value=50)
    path = tmp_path / "sheet.xlsx"
    workbook.save(path)

    assert submitted_by_section_from_workbook(str(path)) == {
        "1.02.01.0693": 36.30,
        "1.02.01.0671": 4.25,
    }


def test_sheet_item_missing_from_boq_is_a_mismatch():
    rows = quantity_mismatches({"99.1": (12.5, 0.0)}, {})
    assert rows[0]["section_number"] == "99.1"
    assert rows[0]["system_calculated"] == 0.0
    assert rows[0]["calculated_diff"] == 12.5
    assert rows[0]["is_non_boq"] is True
