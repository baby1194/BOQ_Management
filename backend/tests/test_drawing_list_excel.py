from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from services.excel_service import ExcelService


def test_export_drawing_list_writes_headers_and_status_labels(tmp_path):
    service = ExcelService(exports_dir=tmp_path)
    rows = [
        SimpleNamespace(
            no=1,
            drawing_type="Hydro",
            planning_office="Office",
            drawing_name="Hotel plan 1",
            cross_sections=None,
            element=None,
            sheet_name="HYD-UT-040",
            edition="0",
            release_date="18.08.2025",
            update_description=None,
            folder_date="01.07.2025",
            file_path=r"C:\Users\user\Documents\plan.pdf",
            notes=None,
            execution_status="to_be_executed",
        )
    ]

    path = Path(service.export_drawing_list(rows, language="en"))
    workbook = load_workbook(path)
    sheet = workbook.active

    assert sheet.title == "List of Drawings"
    assert sheet.cell(1, 1).value == "No"
    assert sheet.cell(1, 4).value == "Drawing Name"
    assert sheet.cell(2, 1).value == 1
    assert sheet.cell(2, 7).value == "HYD-UT-040"
    assert sheet.cell(2, 14).value == "To Be Executed"
    assert sheet.cell(2, 12).hyperlink is not None


def test_export_drawing_list_hebrew_headers(tmp_path):
    service = ExcelService(exports_dir=tmp_path)
    rows = [
        {
            "no": 2,
            "drawing_type": "סוג",
            "planning_office": "",
            "drawing_name": "שם",
            "cross_sections": "",
            "element": "",
            "sheet_name": "A-01",
            "edition": "1",
            "release_date": "01.01.2026",
            "update_description": "",
            "folder_date": "",
            "file_path": "",
            "notes": "",
            "execution_status": "cancelled",
            "created_at": datetime.now().isoformat(),
        }
    ]

    path = Path(service.export_drawing_list(rows, language="he"))
    workbook = load_workbook(path)
    sheet = workbook.active

    assert sheet.title == "רשימת תכניות"
    assert sheet.cell(1, 1).value == "מס׳"
    assert sheet.cell(2, 14).value == "מבוטל"
    assert sheet.sheet_view.rightToLeft is True
