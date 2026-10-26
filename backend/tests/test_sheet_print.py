from pathlib import Path

from openpyxl import Workbook

from services.sheet_print import (
    build_print_grid,
    calculation_worksheet,
    filled_cell_boxes,
    format_cell,
    pdf_beside_workbook,
    print_workbooks,
)


def test_pdf_uses_the_workbook_name_in_the_same_folder():
    workbook = Path("C:/Fatina/1.02/17_1/sheet.xlsx")
    assert pdf_beside_workbook(workbook) == Path("C:/Fatina/1.02/17_1/sheet.pdf")


def test_numbers_have_no_float_noise():
    assert format_cell(13.850000000000001) == "13.85"
    assert format_cell(4.0) == "4"
    assert format_cell(0.1 + 0.2) == "0.3"
    assert format_cell(None) == ""


def test_only_the_filled_block_is_printed():
    book = Workbook()
    ws = book.active
    ws["C3"] = "item"
    ws["E3"] = "qty"
    ws["C5"] = "1.1"
    ws["E5"] = 13.850000000000001
    grid = build_print_grid(ws)
    assert grid.rows == [["item", "qty"], ["1.1", "13.85"]]
    assert grid.right_to_left is False


def test_hebrew_sheet_prints_right_to_left():
    book = Workbook()
    ws = book.active
    ws.sheet_view.rightToLeft = True
    ws["A1"] = "סעיף"
    ws["B1"] = "כמות"
    ws["A2"] = "1.1"
    ws["B2"] = 3
    grid = build_print_grid(ws)
    assert grid.right_to_left is True
    assert grid.rows[0] == ["תומכ", "ףיעס"]
    assert grid.rows[1] == ["3", "1.1"]


def test_print_area_limits_the_block():
    book = Workbook()
    ws = book.active
    ws["A1"] = "keep"
    ws["B1"] = "keep too"
    ws["D1"] = "outside"
    ws.print_area = "A1:B1"
    assert build_print_grid(ws).rows == [["keep", "keep too"]]


def test_merged_title_spans_the_kept_columns():
    book = Workbook()
    ws = book.active
    ws["A1"] = "title"
    ws.merge_cells("A1:C1")
    ws["A2"] = "a"
    ws["B2"] = "b"
    ws["C2"] = "c"
    grid = build_print_grid(ws)
    assert grid.spans == [((0, 0), (2, 0))]


def test_a_long_sheet_repeats_its_first_27_rows_and_keeps_its_own_columns():
    book = Workbook()
    ws = book.active
    ws["A1"] = "title"
    ws["C28"] = "item"
    ws["E28"] = 4
    grid = build_print_grid(ws)
    assert grid.repeat_rows == 27
    assert len(grid.rows) == 28
    assert grid.rows[0] == ["title", "", ""]
    assert grid.rows[1] == ["", "", ""]
    assert grid.rows[27] == ["", "item", "4"]
    assert all(len(row) == 3 for row in grid.rows)


def test_a_short_sheet_does_not_pad_empty_title_rows():
    book = Workbook()
    ws = book.active
    ws["A1"] = "only"
    ws["B3"] = 2
    grid = build_print_grid(ws)
    assert grid.repeat_rows == 0
    assert grid.rows == [["only", ""], ["", "2"]]


def test_borders_are_only_drawn_around_filled_cells():
    boxes = filled_cell_boxes(
        [["title", "", ""], ["", "1", ""]],
        [((0, 0), (1, 0))],
    )
    assert boxes == [((0, 0), (1, 0)), ((1, 1), (1, 1))]


def test_calculation_sheet_is_printed_instead_of_the_first_worksheet():
    book = Workbook()
    book.active.title = "Other"
    book.create_sheet("Calculation")
    assert calculation_worksheet(book).title == "Calculation"


def test_print_writes_a_pdf_next_to_each_workbook(tmp_path):
    folder = tmp_path / "item" / "17_1"
    folder.mkdir(parents=True)
    book = Workbook()
    ws = book.active
    ws.title = "Calculation"
    ws.append(["סעיף", "כמות"])
    ws.append(["1.1", 3])
    book.save(folder / "17_1.xlsx")
    written = print_workbooks([tmp_path], orientation="portrait", margin_mm=0, columns=[1, 2])
    assert written == [str(folder / "17_1.pdf")]
    assert (folder / "17_1.pdf").is_file()
