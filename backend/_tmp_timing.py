import time
from openpyxl import load_workbook

path = r"D:\Ares\___projects___\Johnkh69\Calculation Sheets\Calculation Sheets\12.xlsx"
start = time.perf_counter()
wb = load_workbook(path, read_only=True, data_only=True)
ws = wb["Calculation"]
print("opened", round(time.perf_counter() - start, 3))
print("max_row", ws.max_row, "max_column", ws.max_column)
start = time.perf_counter()
count = 0
for _ in ws.iter_rows(max_col=ws.max_column, values_only=True):
    count += 1
    if count == 5 or count % 5000 == 0:
        print("rows", count, "elapsed", round(time.perf_counter() - start, 3))
    if count >= 20000:
        print("stopping early")
        break
print("iterated", count, "elapsed", round(time.perf_counter() - start, 3))
wb.close()
